import pytest

from document.loader import load_document, DocumentError


def test_txt_utf8(tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("Đoạn một nói về học máy.\n\nĐoạn hai nói về dữ liệu.", encoding="utf-8")
    doc = load_document(str(p))
    assert doc.source_type == "txt" and len(doc.paragraphs) == 2


def test_docx_paragraphs_and_tables(tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("Tiêu đề tài liệu")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Model", "Accuracy"
    t.cell(1, 0).text, t.cell(1, 1).text = "CNN", "0.95"
    p = tmp_path / "a.docx"
    d.save(str(p))
    doc = load_document(str(p))
    assert doc.paragraphs[0] == "Tiêu đề tài liệu"
    assert "Model | Accuracy" in doc.paragraphs and "CNN | 0.95" in doc.paragraphs


def test_pptx_text_notes_and_pages(tmp_path):
    from pptx import Presentation
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Giới thiệu Scrum"
    s.placeholders[1].text = "Sprint kéo dài tối đa một tháng"
    s.notes_slide.notes_text_frame.text = "Ghi chú diễn giả"
    p = tmp_path / "a.pptx"
    prs.save(str(p))
    doc = load_document(str(p))
    assert doc.pages == 1
    assert "Giới thiệu Scrum." in doc.paragraphs
    assert "Ghi chú diễn giả." in doc.paragraphs


def test_unsupported_extension(tmp_path):
    p = tmp_path / "a.xyz"
    p.write_text("x")
    with pytest.raises(DocumentError):
        load_document(str(p))


def test_empty_file(tmp_path):
    p = tmp_path / "a.txt"
    p.write_bytes(b"")
    with pytest.raises(DocumentError):
        load_document(str(p))


def test_corrupted_docx(tmp_path):
    p = tmp_path / "a.docx"
    p.write_bytes(b"day khong phai file docx")
    with pytest.raises(DocumentError):
        load_document(str(p))


def test_image_without_ocr_raises(tmp_path):
    p = tmp_path / "a.png"
    p.write_bytes(b"\x89PNG fake")
    with pytest.raises(DocumentError):
        load_document(str(p), ocr_func=None)


def test_image_with_fake_ocr(tmp_path):
    p = tmp_path / "a.jpg"
    p.write_bytes(b"fake jpg bytes")
    doc = load_document(str(p), ocr_func=lambda data, mime: "Văn bản từ ảnh")
    assert doc.used_ocr and doc.paragraphs == ["Văn bản từ ảnh"]
