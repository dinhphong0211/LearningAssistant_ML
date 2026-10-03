import pytest

from summarization import abstractive
from summarization.abstractive import TransformerSummarizer, SKIP_MARK
from document.loader import load_document
import pipeline
from pipeline import run_pipeline, FULL_LABEL, QUICK_LABEL


class Scripted:
    """Giả lập Gemini: nhận prompt, trả lời theo hàm `behave(chunk_text)`."""
    behave = staticmethod(lambda chunk: chunk)

    def __init__(self, name):
        self.name = name

    def generate_content(self, prompt):
        chunk = prompt.split("Văn bản:\n", 1)[1]
        out = Scripted.behave(chunk)
        if isinstance(out, Exception):
            raise out

        class R:
            text = out
        return R()


@pytest.fixture(autouse=True)
def fake(monkeypatch):
    Scripted.behave = staticmethod(lambda chunk: chunk)
    monkeypatch.setattr(abstractive, "configure_gemini", lambda: None)
    monkeypatch.setattr(abstractive.genai, "GenerativeModel", Scripted)
    monkeypatch.setattr(abstractive.time, "sleep", lambda s: None)


def _summarizer():
    return TransformerSummarizer("m", fallback_models=(), retry_delays=())


def _doc(n=30):
    return "\n\n".join(f"Đoạn {i}. " + "Học máy là một nhánh của AI. " * 25 for i in range(n))


def test_faithful_copy_is_kept_in_order_and_complete():
    text = _doc()
    out, rep = _summarizer().prepare_reading_text(text, chunk_chars=3000)
    assert out.split() == text.split()
    assert rep["ai"] == rep["chunks"] > 1 and rep["kept_raw"] == 0


def test_removal_of_author_line_is_accepted():
    text = "Nguyễn Văn A, Đại học X, a@x.edu\n\n" + _doc(3)
    Scripted.behave = staticmethod(lambda c: c.replace("Nguyễn Văn A, Đại học X, a@x.edu\n\n", ""))
    out, rep = _summarizer().prepare_reading_text(text, chunk_chars=100000)
    assert "Nguyễn Văn A" not in out and "Đoạn 0." in out and rep["ai"] == 1


def test_silent_summarisation_is_rejected_and_original_kept():
    text = _doc(10)
    Scripted.behave = staticmethod(lambda c: "Tài liệu nói về học máy.")   # tóm tắt lén
    out, rep = _summarizer().prepare_reading_text(text, chunk_chars=100000)
    assert out.split() == text.split()
    assert rep["kept_raw"] == 1 and rep["ai"] == 0


def test_api_error_keeps_original_and_is_reported():
    text = _doc(5)
    Scripted.behave = staticmethod(lambda c: ValueError("API key not valid"))
    out, rep = _summarizer().prepare_reading_text(text, chunk_chars=100000)
    assert out.split() == text.split()
    assert rep["kept_raw"] == 1 and "API key" in rep["errors"][0]


def test_noise_only_short_chunk_is_dropped_but_long_is_not():
    toc = "Mục lục\n\nChương 1 ..... 3\n\nChương 2 ..... 9"
    Scripted.behave = staticmethod(lambda c: SKIP_MARK)
    out, rep = _summarizer().prepare_reading_text(toc, chunk_chars=100000)
    assert out == "" and rep["dropped"] == 1


def test_long_chunk_marked_skip_is_kept():
    big = _doc(30)
    assert len(big) > abstractive.SKIP_MAX_CHARS
    Scripted.behave = staticmethod(lambda c: SKIP_MARK)
    out, rep = _summarizer().prepare_reading_text(big, chunk_chars=100000)
    assert out.split() == big.split() and rep["dropped"] == 0 and rep["kept_raw"] == 1


# ---------- pipeline ----------
class FakeSummarizer:
    model = None

    def __init__(self, fn=None, report=None):
        self.fn, self.report = fn, report
        self.calls = 0

    def prepare_reading_text(self, text, chunk_chars=6000):
        self.calls += 1
        out = self.fn(text) if self.fn else text
        return out, self.report or {"chunks": 1, "ai": 1, "dropped": 0, "kept_raw": 0, "errors": []}


@pytest.fixture
def txt(tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("Học máy là gì.\n\n12\n\nTrang 13\n\nĐây là đoạn hai về dữ liệu.", encoding="utf-8")
    return str(p)


def test_full_mode_keeps_all_content_and_drops_page_numbers(txt):
    s = FakeSummarizer()
    r = run_pipeline(txt, s, FULL_LABEL)
    assert r["mode"] == "full" and r["baseline"] == ""
    assert r["summary"] == "Học máy là gì.\n\nĐây là đoạn hai về dữ liệu."
    assert s.calls == 1


def test_full_mode_without_ai_never_calls_gemini(txt):
    s = FakeSummarizer()
    r = run_pipeline(txt, s, FULL_LABEL, use_ai_cleanup=False)
    assert s.calls == 0
    assert "Đây là đoạn hai về dữ liệu." in r["summary"]
    assert any("Chưa dùng Gemini" in w for w in r["warnings"])


def test_full_mode_all_chunks_failed_returns_gemini_error(txt):
    rep = {"chunks": 1, "ai": 0, "dropped": 0, "kept_raw": 1, "errors": ["quota"]}
    r = run_pipeline(txt, FakeSummarizer(report=rep), FULL_LABEL)
    assert r["summary"].startswith(abstractive.GEMINI_ERROR_PREFIX)


def test_full_mode_partial_failure_warns(txt):
    rep = {"chunks": 4, "ai": 3, "dropped": 0, "kept_raw": 1, "errors": ["x"]}
    r = run_pipeline(txt, FakeSummarizer(report=rep), FULL_LABEL)
    assert any("1/4" in w for w in r["warnings"])


def test_full_mode_progress_has_four_steps(txt):
    seen = []
    run_pipeline(txt, FakeSummarizer(), FULL_LABEL, progress=lambda i, t, l: seen.append((i, t)))
    assert seen == [(0, 4), (1, 4), (2, 4), (3, 4)]


def test_full_mode_only_page_numbers_is_error(tmp_path):
    from document.loader import DocumentError
    p = tmp_path / "b.txt"
    p.write_text("12\n\n13", encoding="utf-8")
    with pytest.raises(DocumentError):
        run_pipeline(str(p), FakeSummarizer(), FULL_LABEL)


# ---------- loader: PPTX ----------
def test_pptx_footer_and_slide_number_skipped_only_when_asked(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    layout = prs.slide_layouts[1]
    s = prs.slides.add_slide(layout)
    s.shapes.title.text = "Giới thiệu Scrum"
    s.placeholders[1].text = "Sprint kéo dài tối đa một tháng"
    # chèn ô số slide và chân trang như mẫu slide thật
    from pptx.enum.shapes import PP_PLACEHOLDER as P
    from lxml import etree
    for ph_type, text in (("sldNum", "7"), ("ftr", "Công ty ABC - Nội bộ")):
        sp = etree.fromstring(
            '<p:sp xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
            'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
            f'<p:nvSpPr><p:cNvPr id="{90 if ph_type == "ftr" else 91}" name="{ph_type}"/><p:cNvSpPr/>'
            f'<p:nvPr><p:ph type="{ph_type}" idx="{11 if ph_type == "ftr" else 12}"/></p:nvPr></p:nvSpPr>'
            '<p:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="914400" cy="300000"/></a:xfrm></p:spPr>'
            f'<p:txBody><a:bodyPr/><a:p><a:r><a:t>{text}</a:t></a:r></a:p></p:txBody></p:sp>'
        )
        s.shapes._spTree.append(sp)
    path = str(tmp_path / "a.pptx")
    prs.save(path)

    default = load_document(path).paragraphs
    clean = load_document(path, clean_page_furniture=True).paragraphs
    assert "Công ty ABC - Nội bộ." in default and "7." in default      # hành vi cũ không đổi
    assert "Công ty ABC - Nội bộ." not in clean and "7." not in clean
    assert "Giới thiệu Scrum." in clean and "Sprint kéo dài tối đa một tháng." in clean
