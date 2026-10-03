"""Xuất bản tóm tắt ra .txt và .docx để người dùng tải về."""
import io


def summary_to_txt_bytes(title, summary, terms=None):
    parts = [title, "=" * len(title), "", summary.strip()]
    if terms:
        parts += ["", "Thuật ngữ chính:", ", ".join(terms)]
    return ("\n".join(parts) + "\n").encode("utf-8")


def summary_to_docx_bytes(title, summary, domain=None, terms=None, heading="Nội dung tóm tắt"):
    import docx

    d = docx.Document()
    d.add_heading(title, level=1)
    if domain:
        d.add_paragraph(f"Lĩnh vực: {domain}")
    d.add_heading(heading, level=2)
    for para in [p.strip() for p in summary.split("\n") if p.strip()]:
        d.add_paragraph(para)
    if terms:
        d.add_heading("Thuật ngữ chính", level=2)
        d.add_paragraph(", ".join(terms))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
