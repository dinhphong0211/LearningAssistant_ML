"""Đọc nhiều định dạng tài liệu và trả về cùng một cấu trúc LoadedDocument.

Hỗ trợ: PDF (có text layer, và PDF scan qua OCR), DOCX, PPTX, TXT, PNG/JPG/JPEG.
Mọi lỗi người dùng có thể gặp được ném ra dưới dạng DocumentError với thông báo tiếng Việt.

Một file: load_document(path).
Nhiều file riêng lẻ gộp thành một tài liệu: load_documents([path1, (path2, "ten_hien_thi.pdf"), ...]).
"""
import hashlib
import os
from dataclasses import dataclass, field

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".pptx", ".txt", ".png", ".jpg", ".jpeg")
IMAGE_MIME = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}
MIN_CHARS_PER_PAGE = 50   # dưới ngưỡng này coi như PDF scan (không có text layer)
MAX_FILES = 10            # số file tối đa gộp trong một lần


class DocumentError(Exception):
    """Lỗi có thông báo thân thiện để hiển thị cho người dùng."""


@dataclass
class LoadedDocument:
    paragraphs: list
    source_type: str
    pages: int = None
    used_ocr: bool = False
    warnings: list = field(default_factory=list)
    # Mỗi phần tử: {"name", "type", "pages", "paragraphs", "chars"}; có từ load_documents
    sources: list = field(default_factory=list)


def get_extension(filename):
    return os.path.splitext(filename or "")[1].lower()


def load_document(path, ocr_func=None, max_ocr_pages=10, clean_page_furniture=False):
    """clean_page_furniture=True: bỏ số trang, header/footer lặp lại (PDF, ảnh OCR) và
    ô số trang/chân trang/ngày của slide (PPTX). Mặc định tắt để không đổi hành vi cũ."""
    ext = get_extension(path)
    if ext not in SUPPORTED_EXTENSIONS:
        raise DocumentError(
            f"Định dạng '{ext or '(không có)'}' chưa được hỗ trợ. "
            f"Hỗ trợ: {', '.join(SUPPORTED_EXTENSIONS)}."
        )
    if not os.path.exists(path):
        raise DocumentError("Không tìm thấy file. Vui lòng tải lên lại.")
    if os.path.getsize(path) == 0:
        raise DocumentError("File rỗng (0 byte).")

    try:
        if ext == ".pdf":
            doc = _load_pdf(path, ocr_func, max_ocr_pages, clean_page_furniture)
        elif ext == ".docx":
            doc = _load_docx(path)
        elif ext == ".pptx":
            doc = _load_pptx(path, clean_page_furniture)
        elif ext == ".txt":
            doc = _load_txt(path)
        else:
            doc = _load_image(path, ext, ocr_func)
    except DocumentError:
        raise
    except Exception as e:
        raise DocumentError(
            f"Không đọc được file {ext}. File có thể bị hỏng hoặc có mật khẩu. (Chi tiết: {e})"
        )

    if not "".join(doc.paragraphs).strip():
        raise DocumentError("Không trích xuất được nội dung văn bản từ file này.")
    return doc


# ---------------- NHIỀU FILE ----------------
def _stem(name):
    base = os.path.splitext(os.path.basename(name or ""))[0]
    return " ".join(base.replace("_", " ").replace("-", " ").split()) or "không tên"


def _normalize_items(items):
    """Chấp nhận: một đường dẫn, hoặc danh sách gồm đường dẫn / (đường dẫn, tên hiển thị)."""
    if isinstance(items, (str, os.PathLike)):
        items = [items]
    entries = []
    for it in items or []:
        if isinstance(it, (tuple, list)):
            path, name = it[0], it[1]
        else:
            path, name = it, os.path.basename(os.fspath(it))
        entries.append((os.fspath(path), str(name)))
    return entries


def _source_info(name, doc):
    return {
        "name": name,
        "type": doc.source_type,
        "pages": doc.pages,
        "paragraphs": len(doc.paragraphs),
        "chars": sum(len(p) for p in doc.paragraphs),
    }


def load_documents(items, ocr_func=None, max_ocr_pages=10, clean_page_furniture=False, add_headings=True):
    """Đọc nhiều file riêng lẻ rồi gộp thành MỘT LoadedDocument theo thứ tự đã cho.

    - Một file: kết quả y hệt load_document (không thêm tiêu đề), chỉ có thêm `sources`.
    - Nhiều file: trước nội dung mỗi file thêm một dòng "Tài liệu i: tên file." để người nghe/đọc
      biết đang chuyển sang file khác (tắt bằng add_headings=False).
    - File lỗi (nhiều file): bỏ qua và ghi vào warnings; chỉ báo lỗi khi KHÔNG đọc được file nào.
    - File trùng nội dung với file đã đọc: bỏ qua và ghi vào warnings.
    """
    entries = _normalize_items(items)
    if not entries:
        raise DocumentError("Chưa chọn file nào.")
    if len(entries) > MAX_FILES:
        raise DocumentError(f"Tối đa {MAX_FILES} file mỗi lần (đang chọn {len(entries)}).")

    ok, problems, seen = [], [], {}
    for path, name in entries:
        try:
            doc = load_document(path, ocr_func=ocr_func, max_ocr_pages=max_ocr_pages,
                                clean_page_furniture=clean_page_furniture)
        except DocumentError as e:
            if len(entries) == 1:
                raise
            problems.append(f"Bỏ qua '{name}': {e}")
            continue
        digest = hashlib.sha1("\n".join(doc.paragraphs).encode("utf-8", "ignore")).hexdigest()
        if digest in seen:
            problems.append(f"Bỏ qua '{name}': nội dung trùng với '{seen[digest]}'.")
            continue
        seen[digest] = name
        ok.append((name, doc))

    if not ok:
        raise DocumentError("Không đọc được file nào. " + " ".join(problems))

    if len(ok) == 1:
        name, doc = ok[0]
        doc.sources = [_source_info(name, doc)]
        doc.warnings = list(doc.warnings) + problems
        return doc

    paragraphs, warnings, sources = [], [], []
    pages_total, used_ocr = 0, False
    for i, (name, doc) in enumerate(ok, 1):
        if add_headings:
            paragraphs.append(f"Tài liệu {i}: {_stem(name)}.")
        paragraphs.extend(doc.paragraphs)
        warnings += [f"[{name}] {w}" for w in doc.warnings]
        sources.append(_source_info(name, doc))
        pages_total += doc.pages or 0
        used_ocr = used_ocr or doc.used_ocr
    return LoadedDocument(
        paragraphs,
        f"Gộp {len(ok)} file",
        pages=pages_total or None,
        used_ocr=used_ocr,
        warnings=warnings + problems,
        sources=sources,
    )


# ---------------- PDF ----------------
def _load_pdf(path, ocr_func, max_ocr_pages, clean_page_furniture=False):
    import pymupdf

    pdf = pymupdf.open(path)
    try:
        if pdf.needs_pass:
            raise DocumentError("PDF có mật khẩu. Hãy gỡ mật khẩu rồi tải lại.")
        n_pages = len(pdf)
        pages_blocks = [(page.rect.height, page.get_text("blocks")) for page in pdf]
        if clean_page_furniture:
            from preprocessing.boilerplate import filter_pdf_furniture
            blocks = filter_pdf_furniture(pages_blocks)
        else:
            blocks = [b for _, page_blocks in pages_blocks for b in page_blocks]

        try:
            from preprocessing.cleaner import merge_pdf_blocks
            paragraphs = list(merge_pdf_blocks(blocks))
        except ImportError:
            paragraphs = [b[4].strip() for b in blocks if len(b) > 4 and str(b[4]).strip()]

        total_chars = sum(len(p) for p in paragraphs)
        if total_chars >= MIN_CHARS_PER_PAGE * max(n_pages, 1):
            return LoadedDocument(paragraphs, "pdf", pages=n_pages)

        # Có vẻ là PDF scan: dùng OCR nếu có
        if ocr_func is None:
            raise DocumentError(
                "PDF này gần như không có văn bản (có thể là bản scan) và OCR chưa được bật."
            )
        warnings = []
        pages_to_read = min(n_pages, max_ocr_pages)
        if n_pages > max_ocr_pages:
            warnings.append(
                f"PDF scan có {n_pages} trang, chỉ OCR {max_ocr_pages} trang đầu để tiết kiệm thời gian và quota."
            )
        texts = []
        for i in range(pages_to_read):
            png = pdf[i].get_pixmap(dpi=150).tobytes("png")
            texts.append(ocr_func(png, "image/png"))
        warnings.append("Nội dung lấy bằng OCR (Gemini), công thức và ký hiệu có thể sai.")
        if clean_page_furniture:
            from preprocessing.boilerplate import strip_repeated_page_lines
            texts = strip_repeated_page_lines(texts)
        paragraphs = [t for t in texts if t.strip()]
        return LoadedDocument(paragraphs, "pdf (scan)", pages=n_pages, used_ocr=True, warnings=warnings)
    finally:
        pdf.close()


# ---------------- DOCX ----------------
def _load_docx(path):
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(path)
    paragraphs = []
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, d).text.strip()
            if text:
                paragraphs.append(text)
        elif tag == "tbl":
            for row in Table(child, d).rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    paragraphs.append(" | ".join(cells))
    return LoadedDocument(paragraphs, "docx")


# ---------------- PPTX ----------------
def _end_sentence(text):
    text = text.strip()
    return text if text.endswith((".", "!", "?", ":", ";")) else text + "."


def _is_furniture_placeholder(shape):
    """Ô số slide / chân trang / ngày / đầu trang do mẫu slide sinh ra."""
    try:
        if not shape.is_placeholder:
            return False
        from pptx.enum.shapes import PP_PLACEHOLDER as P

        bad = {getattr(P, n) for n in ("SLIDE_NUMBER", "FOOTER", "DATE", "HEADER") if hasattr(P, n)}
        return shape.placeholder_format.type in bad
    except Exception:
        return False


def _load_pptx(path, clean_page_furniture=False):
    from pptx import Presentation

    prs = Presentation(path)
    paragraphs = []
    for i, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            if clean_page_furniture and _is_furniture_placeholder(shape):
                continue
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    t = "".join(r.text for r in p.runs).strip()
                    if t:
                        paragraphs.append(_end_sentence(t))
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    cells = [c.text.strip() for c in row.cells]
                    if any(cells):
                        paragraphs.append(_end_sentence(" | ".join(cells)))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                paragraphs.append(_end_sentence(notes))
    return LoadedDocument(paragraphs, "pptx", pages=len(prs.slides))


# ---------------- TXT ----------------
def _load_txt(path):
    raw = open(path, "rb").read()
    text = None
    for enc in ("utf-8-sig", "utf-16", "cp1258", "cp1252"):
        try:
            text = raw.decode(enc)
            break
        except (UnicodeDecodeError, UnicodeError):
            continue
    if text is None:
        raise DocumentError("Không xác định được bảng mã của file .txt (thử lưu lại dưới dạng UTF-8).")
    paragraphs = [p.strip() for p in text.replace("\r\n", "\n").split("\n\n") if p.strip()]
    return LoadedDocument(paragraphs, "txt")


# ---------------- ẢNH ----------------
def _load_image(path, ext, ocr_func):
    if ocr_func is None:
        raise DocumentError("Cần bật OCR (Gemini) để đọc file ảnh.")
    with open(path, "rb") as f:
        data = f.read()
    text = ocr_func(data, IMAGE_MIME[ext])
    return LoadedDocument(
        [text] if text.strip() else [],
        "image",
        used_ocr=True,
        warnings=["Nội dung lấy bằng OCR (Gemini), công thức và ký hiệu có thể sai."],
    )
